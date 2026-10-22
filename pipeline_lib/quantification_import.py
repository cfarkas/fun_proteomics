"""Read provider workbooks without guessing a statistical-results sheet as quantities."""
from pathlib import Path
import re
import pandas as pd

IDENTIFIERS = ('Protein.Group', 'Protein', 'FeatureID', 'Protein ID', 'protein_group', 'group_key',
               'Protein IDs', 'Protein.IDs', 'Majority protein IDs', 'Uniprot', 'Fasta headers', 'Fasta.headers')


def header_score(header):
    """Prefer explicit groups or annotated FASTA headers over accession-only copies."""
    fields = set(header)
    if not fields.intersection(IDENTIFIERS): return 0
    measurements = any(re.match(r'^(?:LFQ intensity |LFQ\.intensity\.|Intensity |Intensity\.)', c)
                       or c.endswith((' MaxLFQ Intensity', ' Intensity', ' Spectral Count')) for c in fields)
    if not measurements and len(fields) < 2: return 0
    # Statistical summary sheets are never quantity matrices.
    if not measurements and any('p.val' in c.lower() or 'fold change' in c.lower() for c in fields): return 0
    return 30 if fields.intersection(IDENTIFIERS[:9]) else 20 if fields.intersection({'Fasta headers','Fasta.headers'}) else 10


def workbook_table(path, sheet=None):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet and sheet not in book.sheetnames:
            raise ValueError('Excel worksheet not found: '+sheet+'. Available: '+', '.join(book.sheetnames))
        candidates = []
        for name in [sheet] if sheet else book.sheetnames:
            for number, row in enumerate(book[name].iter_rows(max_row=25, values_only=True), 1):
                header = [str(v) if v is not None else '' for v in row]
                while header and not header[-1]: header.pop()
                score = header_score(header)
                if score:
                    candidates.append((score, name, number, header)); break
        if not candidates:
            raise ValueError('No protein quantity worksheet found. Select a sheet with protein identifiers and sample intensities; statistical result tables are not intensities.')
        best = max(c[0] for c in candidates); candidates = [c for c in candidates if c[0] == best]
        if len(candidates) != 1:
            raise ValueError('Multiple protein quantity worksheets: '+', '.join(c[1] for c in candidates)+'. Select one with --sheet / Excel worksheet.')
        _, name, number, header = candidates[0]
        if len(set(header)) != len(header) or any(not c.strip() for c in header):
            raise ValueError('Quantification table has empty or duplicate column names')
        values = list(book[name].iter_rows(min_row=number+1, max_col=len(header), values_only=True))
        frame = pd.DataFrame(values, columns=header).dropna(how='all').fillna('').astype(str)
        frame.attrs['import'] = dict(format='excel', worksheet=name, header_row=number)
        return frame
    finally:
        book.close()


def adapt_table(table):
    """Normalize MaxQuant names; retain reported annotations and incomplete group flags."""
    table = table.copy(); audit = dict(table.attrs.get('import', {}))
    rename = {}
    for column in table:
        for prefix, suffix in [('LFQ intensity ', ' MaxLFQ Intensity'), ('LFQ.intensity.', ' MaxLFQ Intensity'),
                               ('Intensity ', ' Intensity'), ('Intensity.', ' Intensity')]:
            if column.startswith(prefix) and column[len(prefix):]:
                rename[column] = column[len(prefix):]+suffix; break
    names = [rename.get(c,c) for c in table]
    if len(set(names)) != len(names): raise ValueError('Duplicate sample columns after quantity-name normalization')
    originals = {rename.get(c,c):c for c in table}; table = table.rename(columns=rename)
    aliases = {'Protein.Group': ('Protein IDs','Protein.IDs','Majority protein IDs','Uniprot'),
               'Genes': ('Gene names','Gene.names','Gene Name'),
               'Description': ('Protein names','Protein.names'), 'Fasta headers': ('Fasta.headers',)}
    for dest, sources in aliases.items():
        if dest not in table:
            source = next((c for c in sources if c in table), None)
            if source: table[dest] = table[source]
    header_ids = False
    if 'Fasta headers' in table:
        def parse(value):
            parts = [p.strip() for p in str(value).split(';') if p.strip()]
            identifiers = []; genes = []; names = []
            complete = bool(parts)
            for part in parts:
                match = re.match(r'(?:(?:REV|CON)__)?(?:sp|tr)\|([^|\s]+)\|([^\s|]+)', part)
                if match:
                    identifiers.append(match.group(1)); names.append(match.group(2))
                else: complete = False
                symbol = re.search(r'(?:^|\s)GN=([^\s;]+)', part)
                if symbol: genes.append(symbol.group(1))
                # MaxQuant service exports sometimes clip long headers to 256 characters.
                if not re.search(r'\sSV=\d+\s*$', part): complete = False
            return ';'.join(sorted(set(identifiers))), ';'.join(sorted(set(genes))), ';'.join(sorted(set(names))), complete
        parsed = [parse(v) for v in table['Fasta headers']]
        header_ids = not any(c in table for c in IDENTIFIERS[:-2])
        if header_ids: table['Protein.Group'] = [v[0] for v in parsed]
        if 'Genes' not in table: table['Genes'] = [v[1] for v in parsed]
        if 'Protein.Names' not in table: table['Protein.Names'] = [v[2] for v in parsed]
        table['reported_fasta_headers'] = table['Fasta headers']
        if header_ids:
            table['protein_group_membership'] = ['complete' if v[3] else 'partial_header' for v in parsed]
            audit['incomplete_protein_headers'] = sum(not v[3] for v in parsed)
    flags = pd.Series(False, index=table.index)
    for col in ('Reverse','Potential contaminant','Potential.contaminant','Only identified by site','Only.identified.by.site'):
        if col in table: flags |= table[col].str.strip().str.lower().isin(['+','true','1'])
    if 'Fasta headers' in table:
        flags |= table['Fasta headers'].str.contains(r'(?:^|;)(?:REV__|CON__|rev_|contam_)', regex=True)
    audit['excluded_provider_flag_rows'] = int(flags.sum())
    table = table.loc[~flags].copy()
    if header_ids:
        # Some service workbooks repeat the same group as an entirely empty row.
        # Remove only zero/missing duplicates; never combine observed quantities.
        measures=[c for c in table if c.endswith(' MaxLFQ Intensity')] or [c for c in table if c.endswith(' Intensity')]
        if measures:
            numeric=table[measures].replace({'':0,'NA':0,'NaN':0,'nan':0,'N/A':0}).apply(pd.to_numeric,errors='coerce')
            empty=numeric.eq(0).all(axis=1); removed=[]
            for key, group in table.groupby('Protein.Group',sort=False):
                if len(group)<2: continue
                observed=[i for i in group.index if not empty.loc[i]]
                keep=observed[0] if observed else group.index[0]
                for i in group.index:
                    if i!=keep and empty.loc[i]: removed.append({'source_data_row':int(i)+1,'protein_group':key,'reason':'duplicate group with no measured quantities'})
            if removed: table=table.drop(index=[v['source_data_row']-1 for v in removed])
            audit['removed_empty_duplicates']=removed
    table.attrs.update({'source_columns': originals, 'import': audit})
    return table


def read_quantification(path, sheet=None):
    from quantification import read_table
    with Path(path).open('rb') as handle: excel = handle.read(4).startswith(b'PK')
    if sheet and not excel: raise ValueError('--sheet applies only to Excel workbooks')
    return adapt_table(workbook_table(path, sheet) if excel else read_table(path))
