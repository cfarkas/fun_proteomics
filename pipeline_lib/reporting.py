#!/usr/bin/env python3
"""Export explicit protein-group counts and abundance matrices; no imputation."""
import argparse
import csv
import html
import json
import math
from pathlib import Path
from openpyxl import Workbook
from ptm_counts import export as export_ptm_counts
from pipeline import read, save, selected_samples, digest
from immunopeptidomics import enabled as hla_enabled, FDR_DEFINITION

ANNOTATIONS = ['Protein Group', 'Protein', 'Protein ID', 'Gene', 'Description', 'Indistinguishable Proteins']
DEFINITIONS = {
    'unit': 'One nondecoy, noncontaminant combined_protein.tsv row (protein group), restricted to configured organism.',
    'direct_protein_groups': 'Protein groups with sample Spectral Count > 0.',
    'maxlfq_protein_groups': 'Protein groups with sample MaxLFQ Intensity > 0. Blank means unavailable, not zero.',
    'intensity_protein_groups': 'Protein groups with ordinary sample Intensity > 0; not interchangeable with MaxLFQ.',
    'maxlfq_without_direct_ms2': 'Positive MaxLFQ with zero sample spectral count: evidence compatible with MBR; not proof of transfer for every row.',
    'matrices': 'Unlogged, unnormalized by this exporter, not imputed; source zero/blank values retained. IonQuant normalization is enabled upstream.',
    'fdr': 'Percolator rescoring, ProteinProphet inference, Philosopher sequential picked 1% protein FDR; IonQuant ionfdr=0.01, peptidefdr=1, proteinfdr=1 as workflow settings.',
    'comparison': 'Each cohort and search mode is inferred/filtered separately. Count differences are reported; no union at an asserted joint FDR.',
}


def definitions(plan):
    result = dict(DEFINITIONS)
    if hla_enabled(plan):
        result.update(fdr=FDR_DEFINITION,
                      unit='Supporting source-protein groups from HLA-enriched peptides; not source-protein expression or protein-FDR-controlled identifications.',
                      maxlfq_without_direct_ms2='Positive MaxLFQ without direct MS2; MBR is disabled in this HLA workflow.',
                      primary='immunopeptidomics/peptides.csv and peptide_spectral_counts.csv, derived from filtered PSMs independently of protein groups.')
    return result


def table(path):
    with Path(path).open(newline='') as handle:
        reader = csv.DictReader(handle, delimiter='\t')
        return reader.fieldnames or [], list(reader)


def write_table(path, rows, fields=None):
    fields = fields or list(rows[0])
    with Path(path).open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fields, delimiter='\t')
        writer.writeheader()
        writer.writerows(rows)


def number(value):
    if value in ('', None, 'NA', 'NaN', 'nan'):
        return 0.0
    result = float(value)
    if not math.isfinite(result) or result < 0:
        raise ValueError(f'Invalid negative or nonfinite measurement: {value}')
    return result


def workbook(path, sheets):
    book = Workbook()
    book.remove(book.active)
    for name, fields, rows in sheets:
        sheet = book.create_sheet(name[:31])
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(k, '') for k in fields])
            # Treat annotation strings literally, even when they begin with =.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    book.save(path)


def summarize(fields, rows, samples, cohort, mode, organism):
    retained = [r for r in rows if not r['Protein'].startswith(('rev_', 'contam_'))
                and (organism is None or r.get('Organism') == organism)]
    if len({r['Protein'] for r in retained}) != len(retained):
        raise ValueError('Duplicate lead Protein identifiers; inspect the source grouping before export')
    metrics = {'spectral_counts': 'Spectral Count', 'maxlfq_intensity': 'MaxLFQ Intensity', 'intensity': 'Intensity'}
    available, matrices = {}, {}
    for metric, suffix in metrics.items():
        present = [s['sample_id'] + ' ' + suffix in fields for s in samples]
        if any(present) and not all(present):
            raise ValueError(f'Incomplete {suffix} sample columns')
        if metric == 'spectral_counts' and not all(present):
            raise ValueError('Missing exact sample Spectral Count columns; check manifest sample IDs')
        if metric == 'maxlfq_intensity' and len(samples) > 1 and not all(present):
            raise ValueError('A multi-run cohort is missing MaxLFQ columns; quantification may be incomplete')
        available[metric] = all(present)
        if not all(present):
            continue
        matrix = []
        for index, row in enumerate(retained, 1):
            entry = {k: row.get(k, '') for k in ANNOTATIONS}
            entry['Protein Group'] = f'PG{index:06d}'
            for sample in samples:
                value = row[sample['sample_id'] + ' ' + suffix]
                numeric = number(value)
                if metric == 'spectral_counts' and numeric != int(numeric):
                    raise ValueError('Noninteger spectral count')
                entry[sample['sample_id']] = '' if value in ('', None, 'NA', 'NaN', 'nan') else (int(numeric) if metric == 'spectral_counts' else numeric)
            matrix.append(entry)
        matrices[metric] = matrix
    counts = []
    for sample in samples:
        sid = sample['sample_id']
        direct = [number(r[sid]) > 0 for r in matrices['spectral_counts']]
        lfq = [number(r[sid]) > 0 for r in matrices.get('maxlfq_intensity', [])]
        counts.append({'cohort': cohort, 'mode': mode, 'sample_id': sid, 'condition': sample['condition'],
            'replicate': sample['replicate'], 'direct_protein_groups': sum(direct),
            'maxlfq_available': available['maxlfq_intensity'],
            'maxlfq_protein_groups': sum(lfq) if available['maxlfq_intensity'] else '',
            'maxlfq_with_direct_ms2': sum(a and b for a, b in zip(direct, lfq)) if available['maxlfq_intensity'] else '',
            'maxlfq_without_direct_ms2': sum(b and not a for a, b in zip(direct, lfq)) if available['maxlfq_intensity'] else '',
            'direct_without_maxlfq': sum(a and not b for a, b in zip(direct, lfq)) if available['maxlfq_intensity'] else '',
            'intensity_protein_groups': sum(number(r[sid]) > 0 for r in matrices['intensity']) if available['intensity'] else ''})
    cohort_rows = [{'cohort': cohort, 'mode': mode, 'samples': len(samples), 'organism': organism or 'all',
        'groups_in_combined_table': len(retained),
        'groups_with_any_direct_ms2': sum(any(number(r[s['sample_id']]) > 0 for s in samples) for r in matrices['spectral_counts']),
        'groups_with_any_maxlfq': sum(any(number(r[s['sample_id']]) > 0 for s in samples) for r in matrices['maxlfq_intensity']) if available['maxlfq_intensity'] else ''}]
    return counts, cohort_rows, matrices


def report(args):
    plan, run = read(args.plan), Path(args.run)
    if not (run / 'SEARCH_COMPLETE').is_file():
        raise ValueError('No successful search completion marker; refusing partial results')
    samples = selected_samples(plan, args.cohort)
    fields, rows = table(run / 'combined_protein.tsv')
    if 'Protein' not in fields or (plan.get('organism') and 'Organism' not in fields):
        raise ValueError('Missing protein/organism annotations')
    counts, cohort_rows, matrices = summarize(fields, rows, samples, args.cohort, args.mode, plan.get('organism'))
    out = Path(args.out)
    out.mkdir(exist_ok=True)
    sheets = []
    for name, data in [('sample_detection_counts', counts), ('cohort_detection_counts', cohort_rows)]:
        write_table(out / (name + '.tsv'), data)
        sheets.append((name, list(data[0]), data))
    matrix_fields = ANNOTATIONS + [s['sample_id'] for s in samples]
    for name, data in matrices.items():
        write_table(out / ('protein_' + name + '.tsv'), data, matrix_fields)
        sheets.append(('protein_' + name, matrix_fields, data))
    provenance = read(run / 'run_provenance.json')
    if not provenance.get('stub'):
        export_ptm_counts(plan, run, args.cohort, out / 'ptm_counts')
        if hla_enabled(plan):
            from immunopeptidomics import export as export_hla
            export_hla(plan, run, args.cohort, out/'immunopeptidomics')
        else:
            from final_results import peptide_counts
            peptide_proteins={r['Protein'] for r in matrices['spectral_counts']}
            peptide_proteins.update(x.strip() for r in matrices['spectral_counts'] for x in r.get('Indistinguishable Proteins','').split(',') if x.strip())
            peptide_counts(run,samples,peptide_proteins,out/'peptide_spectral_counts.csv')
        from optional_analyses import export_isoforms
        options = plan.get('optional_analyses', {})
        if options.get('peptide_isoforms'):
            export_isoforms(plan, out, args.cohort, out/'optional_analyses/peptide_isoforms')
    report_definitions = definitions(plan)
    save(out / 'definitions.json', report_definitions)
    save(out / 'report_provenance.json', {**provenance, 'source_table_sha256': digest(run / 'combined_protein.tsv')})
    sheets.append(('definitions', ['metric', 'definition'], [{'metric': k, 'definition': v} for k, v in report_definitions.items()]))
    workbook(out / 'count_tables.xlsx', sheets)
    if provenance.get('stub'):
        (out / 'STUB_NOT_SCIENTIFIC_RESULTS').write_text('Synthetic test output.\n')


def finalize(args):
    sample_rows, cohort_rows, provenance = [], [], []
    for report_dir in args.reports:
        report_dir = Path(report_dir)
        sample_rows.extend(table(report_dir / 'sample_detection_counts.tsv')[1])
        cohort_rows.extend(table(report_dir / 'cohort_detection_counts.tsv')[1])
        provenance.append(read(report_dir / 'report_provenance.json'))
    sample_rows.sort(key=lambda r: (r['cohort'], r['sample_id'], r['mode']))
    cohort_rows.sort(key=lambda r: (r['cohort'], r['mode']))
    plan = read(args.plan)
    report_definitions = definitions(plan)
    expected = {(s['cohort'], s['sample_id'], mode) for s in plan['samples'] for mode in plan['modes']}
    actual = [(r['cohort'], r['sample_id'], r['mode']) for r in sample_rows]
    if set(actual) != expected or len(actual) != len(expected):
        raise ValueError('Missing or duplicated sample/mode reports; refusing a final count table')
    comparisons = []
    by_sample = {}
    for row in sample_rows:
        by_sample.setdefault((row['cohort'], row['sample_id']), {})[row['mode']] = row
    comparison_fields = ['cohort', 'sample_id', 'condition']
    metrics = ['direct_protein_groups', 'maxlfq_protein_groups', 'maxlfq_without_direct_ms2']
    for metric in metrics:
        comparison_fields += [metric + '_dda', metric + '_ddaplus', metric + '_delta']
    for (cohort, sid), modes in sorted(by_sample.items()):
        if set(modes) != {'DDA', 'DDA+'}:
            continue
        row = {'cohort': cohort, 'sample_id': sid, 'condition': modes['DDA']['condition']}
        for metric in metrics:
            a, b = modes['DDA'][metric], modes['DDA+'][metric]
            row.update({metric + '_dda': a, metric + '_ddaplus': b,
                        metric + '_delta': int(b) - int(a) if a != '' and b != '' else ''})
        comparisons.append(row)
    out = Path('final')
    out.mkdir(exist_ok=True)
    write_table(out / 'sample_detection_counts.tsv', sample_rows)
    write_table(out / 'cohort_detection_counts.tsv', cohort_rows)
    write_table(out / 'dda_vs_ddaplus.tsv', comparisons, comparison_fields)
    # Convert summary count cells back to numbers for usable Excel tables.
    for rows in [sample_rows, cohort_rows, comparisons]:
        for row in rows:
            for key, value in row.items():
                if key not in {'cohort', 'mode', 'sample_id', 'condition', 'replicate', 'organism'} and isinstance(value, str) and value.isdigit():
                    row[key] = int(value)
    workbook(out / 'updated_count_tables.xlsx', [('samples', list(sample_rows[0]), sample_rows),
        ('cohorts', list(cohort_rows[0]), cohort_rows), ('DDA_vs_DDAplus', comparison_fields, comparisons),
        ('definitions', ['metric', 'definition'], [{'metric': k, 'definition': v} for k, v in report_definitions.items()])])
    stub = any(p.get('stub') for p in provenance)
    save(out / 'provenance.json', {'stub': stub, 'runs': provenance, 'definitions': report_definitions})
    title = 'SYNTHETIC TEST — NOT SCIENTIFIC RESULTS' if stub else 'DDA-PASEF detection counts'
    columns = ['cohort', 'mode', 'sample_id', 'direct_protein_groups', 'maxlfq_protein_groups']
    headers = ''.join('<th>' + html.escape(k) + '</th>' for k in columns)
    cells = ''.join('<tr>' + ''.join('<td>' + html.escape(str(r[k])) + '</td>' for k in columns) + '</tr>' for r in sample_rows)
    (out / 'report.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>' + title + '</title>'
        '<style>body{font:16px system-ui;margin:3em;max-width:1100px}th,td{padding:.5em;border-bottom:1px solid #ddd;text-align:left}</style>'
        '<h1>' + title + '</h1><p>Direct MS/MS evidence and positive MaxLFQ are reported separately. Blank MaxLFQ is unavailable.</p>'
        '<p><a href="updated_count_tables.xlsx">Excel count tables</a> · <a href="dda_vs_ddaplus.tsv">Paired count differences</a></p>'
        '<table><tr>' + headers + '</tr>' + cells + '</table><p>Protein-group matrices and complete FragPipe outputs are in each cohort/mode directory.</p></html>')
    (out / ('STUB_COMPLETE_NOT_SCIENTIFIC_RESULTS' if stub else 'PIPELINE_COMPLETE')).write_text('All configured sample/mode reports present.\n')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(required=True)
    p = sub.add_parser('report')
    for name in ['plan', 'run', 'cohort', 'mode', 'out']:
        p.add_argument('--' + name, required=True)
    p.set_defaults(func=report)
    p = sub.add_parser('finalize')
    p.add_argument('--plan', required=True)
    p.add_argument('--reports', nargs='+', required=True)
    p.set_defaults(func=finalize)
    args = parser.parse_args()
    args.func(args)


if __name__ == '__main__':
    main()
